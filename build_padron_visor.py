# -*- coding: utf-8 -*-
"""Construye el visor a partir de padron_refcat + Habitantes + coordenadas existentes.

- Unifica plurifamiliares por parcela (REFCAT_14).
- Une personas del Excel por dirección normalizada.
- Escribe agregados de sexo/edad en cada vivienda para análisis.
"""

from __future__ import annotations

import csv
import json
import math
import re
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
PADRON = DATA / "padron_refcat.csv"
HABITANTES = ROOT / "Habitantes a fecha 08_10_26.xlsx"
GEO_IN = DATA / "viviendas.geojson"
CACHE = DATA / "coords_cache.json"
MANUAL_XY = DATA / "manual_xy.json"
PORTALES_XY_CACHE = DATA / "portales_xy_cache.json"
PORTALES_URL = (
    "https://sit.rivasciudad.es/server/rest/services/"
    "ACCESOS_PORTALES/FeatureServer/0/query"
)
TODAY = date(2026, 10, 18)
AGE_BANDS = [(0, 14), (15, 29), (30, 44), (45, 64), (65, 200)]
AGE_LABELS = ["0-14", "15-29", "30-44", "45-64", "65+"]
VIA_PREFIX = re.compile(
    r"^(CALLE|CL|AVDA|AVENIDA|AV|PASEO|PS|PLAZA|PZ|CAMINO|CM|CTRA|CARRETERA|CR|"
    r"GLORIETA|GT|TRAVESIA|TRVA|RONDA|RD)\s+"
)


def strip_accents(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def norm_key(text: str) -> str:
    text = strip_accents(text).upper()
    text = text.replace("º", " ").replace("ª", " ").replace(".", " ")
    text = re.sub(r"\s*\([^)]*\)", " ", text)
    text = re.sub(r"[^\w\s/-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def pad_num(num: str) -> str:
    digits = re.sub(r"\D", "", num or "")
    if not digits:
        return ""
    return digits.zfill(4)


def street_num_keys(via_numero: str = "", direccion: str = "") -> list[str]:
    """Claves CALLE|numero para cruzar con portales municipales (CALLE=GOYA, ACCESO=21)."""
    keys: list[str] = []
    for raw in (via_numero, direccion):
        text = norm_key(raw)
        if not text:
            continue
        text = VIA_PREFIX.sub("", text)
        m = re.match(r"^(.+?)\s+(\d+)([A-Z]?)$", text)
        if not m:
            # «CALLE X, 0021 …»
            m = re.match(r"^(.+?),\s*0*(\d+)([A-Z]?)", text)
        if not m:
            continue
        street = norm_key(m.group(1))
        street = VIA_PREFIX.sub("", street)
        try:
            num = str(int(m.group(2)))
        except ValueError:
            continue
        keys.append(f"{street}|{num}")
        # sin artículos finales / (DE LAS)
        short = re.sub(r"\b(DE|DEL|LA|LAS|LOS|EL)\b", " ", street)
        short = re.sub(r"\s+", " ", short).strip()
        if short and short != street:
            keys.append(f"{short}|{num}")
    # únicos preservando orden
    seen = set()
    out = []
    for key in keys:
        if key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def load_portales_xy() -> dict[str, tuple[float, float]]:
    """Índice street|numero → (lon, lat) desde ACCESOS_PORTALES (caché local)."""
    if PORTALES_XY_CACHE.exists():
        raw = json.loads(PORTALES_XY_CACHE.read_text(encoding="utf-8"))
        out = {}
        for key, value in (raw or {}).items():
            if value and len(value) >= 2:
                out[key] = (float(value[0]), float(value[1]))
        if out:
            print(f"Portales XY (caché): {len(out)}", flush=True)
            return out

    print("Descargando coordenadas de portales municipales…", flush=True)
    by_key: dict[str, list[tuple[float, float]]] = defaultdict(list)
    offset = 0
    page = 2000
    while True:
        params = {
            "where": "REFCATASTRAL IS NOT NULL AND REFCATASTRAL <> 'SIN UBICACION PRECISA'",
            "outFields": "REFCATASTRAL,ACCESO,CALLE",
            "returnGeometry": "true",
            "outSR": "4326",
            "resultOffset": offset,
            "resultRecordCount": page,
            "f": "json",
        }
        req = urllib.request.Request(
            PORTALES_URL + "?" + urllib.parse.urlencode(params),
            headers={"User-Agent": "Mozilla/5.0"},
        )
        with urllib.request.urlopen(req, timeout=120) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        if data.get("error"):
            raise RuntimeError(data["error"])
        batch = data.get("features") or []
        for feat in batch:
            attrs = feat.get("attributes") or {}
            geom = feat.get("geometry") or {}
            if "x" not in geom or "y" not in geom:
                continue
            calle = norm_key(attrs.get("CALLE") or "")
            calle = VIA_PREFIX.sub("", calle)
            acceso = re.sub(r"\D", "", str(attrs.get("ACCESO") or ""))
            if not calle or not acceso:
                continue
            try:
                num = str(int(acceso))
            except ValueError:
                continue
            lon, lat = float(geom["x"]), float(geom["y"])
            by_key[f"{calle}|{num}"].append((lon, lat))
            rc = (attrs.get("REFCATASTRAL") or "").strip().upper()[:14]
            if len(rc) >= 14:
                by_key[f"{rc}|{calle}|{num}"].append((lon, lat))
        print(f"  {offset + len(batch)}", flush=True)
        if not data.get("exceededTransferLimit") or not batch:
            break
        offset += len(batch)

    payload = {}
    for key, pts in by_key.items():
        lon = sum(p[0] for p in pts) / len(pts)
        lat = sum(p[1] for p in pts) / len(pts)
        payload[key] = [round(lon, 7), round(lat, 7)]
    PORTALES_XY_CACHE.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(f"Portales XY: {len(payload)} claves", flush=True)
    return {k: (v[0], v[1]) for k, v in payload.items()}


def is_portal_level_key(key: str) -> bool:
    """True si la clave es solo calle+número (sin planta/puerta)."""
    k = norm_key(key)
    if not k:
        return False
    if re.search(r"\bP\d+\b", k):
        return False
    if re.search(r"\b(PBJ|P00|PB|BAJO|BJ)\b", k):
        return False
    # «CALLE X 12» / «CALLE X, 0012» / «CALLE X, 0012-A»
    if re.search(r",\s*\d+[A-Z]?(?:-[A-Z0-9]+)?$", k):
        return True
    if re.search(r"\s\d+[A-Z]?$", k):
        return True
    return False


def key_specificity(key: str) -> int:
    """Mayor = más concreta (planta/puerta)."""
    k = norm_key(key)
    score = len(k)
    if re.search(r"\bP\d+\b", k):
        score += 200
    if re.search(r"\bP\d+\s+\S+", k):
        score += 80
    if re.search(r"\b(PBJ|BAJO|BJ)\b", k):
        score += 120
    if is_portal_level_key(k):
        score -= 150
    return score


def address_keys(
    direccion: str,
    via_numero: str = "",
    portal: str = "",
    piso: str = "",
    puerta: str = "",
) -> list[str]:
    """Claves posibles Excel ↔ padrón, de más específica a más genérica."""
    keys: list[str] = []

    def add(value: str):
        value = norm_key(value)
        if value and value not in keys:
            keys.append(value)

    raw = norm_key(direccion or "")
    if raw:
        add(raw)
        add(re.sub(r"\b0+(\d)", r"\1", raw))
        # CALLE X, 0013 P02 → CALLE X 13 P02
        m = re.match(r"^(.+?),\s*0*(\d+)(.*)$", raw)
        if m:
            add(f"{m.group(1)} {m.group(2)}{m.group(3)}")
            add(f"{m.group(1)}, {m.group(2).zfill(4)}{m.group(3)}")

    via = norm_key(via_numero or "")
    piso_u = (piso or "").strip().upper()
    puerta_u = str(puerta or "").strip().upper()
    if via:
        m = re.match(r"^(.+?)\s+(\d+)$", via)
        if m:
            street, num = m.group(1), m.group(2)
            base = f"{street}, {num.zfill(4)}"
            if portal:
                base_portal = f"{street}, {num.zfill(4)}-{portal}"
            else:
                base_portal = base
            extra = []
            if piso_u:
                extra.append(piso_u)
            if puerta_u:
                extra.append(puerta_u)
            if extra:
                add(base_portal + " " + " ".join(extra))
                add(f"{street} {int(num)} " + " ".join(extra))
            # claves de portal solo si la vivienda no trae planta/puerta
            if not piso_u and not puerta_u:
                add(via)
                add(base)
                if portal:
                    add(base_portal)
        elif not piso_u and not puerta_u:
            add(via)

    keys.sort(key=key_specificity, reverse=True)
    return keys


def person_sig(person: dict) -> tuple:
    return (
        person.get("dni") or person.get("nombre") or "",
        person.get("anio"),
        person.get("sexo") or "",
    )


def parse_birth(value) -> date | None:
    if value is None or value == "" or value == "-":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def age_of(born: date | None) -> int | None:
    if not born:
        return None
    years = TODAY.year - born.year
    if (TODAY.month, TODAY.day) < (born.month, born.day):
        years -= 1
    return max(0, years)


def band_index(age: int | None) -> int | None:
    if age is None:
        return None
    for i, (lo, hi) in enumerate(AGE_BANDS):
        if lo <= age <= hi:
            return i
    return None


def planta_txt(piso: str, unica: bool) -> str:
    p = (piso or "").strip().upper()
    if unica and not p:
        return "Unifamiliar"
    if not p:
        return "Sin planta"
    if p in {"PBJ", "P00", "PB", "BAJO", "BJ", "0", "P0", "00"}:
        return "Baja"
    if p.startswith("P") and p[1:].isdigit():
        n = int(p[1:])
        return "Baja" if n == 0 else f"{n}.ª"
    if p.isdigit():
        n = int(p)
        return "Baja" if n == 0 else f"{n}.ª"
    return p


def z_for(piso: str) -> float:
    p = (piso or "").strip().upper()
    if p in {"PBJ", "P00", "PB", "BAJO", "BJ", "0", "P0", "00"}:
        return 1.5
    if p.startswith("P") and p[1:].isdigit():
        return 1.5 + int(p[1:]) * 3.0
    if p.isdigit():
        return 1.5 + int(p) * 3.0
    return 1.5


def load_people_by_address() -> dict[str, list[dict]]:
    import openpyxl

    print("Leyendo habitantes…", flush=True)
    wb = openpyxl.load_workbook(HABITANTES, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    by: dict[str, list[dict]] = defaultdict(list)
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        estado = str(row[idx["Estado"]] or "")
        if estado and estado.lower() not in {"empadronado", "alta"}:
            # seguir incluyendo Empadronado; filtrar bajas evidentes
            if "baja" in estado.lower() or "fallec" in estado.lower():
                continue
        direccion = str(row[idx["Dirección"]] or "").strip()
        if not direccion:
            continue
        born = parse_birth(row[idx["Fecha nacimiento"]])
        if born is None:
            year = row[idx.get("año nacimiento")]
            try:
                born = date(int(year), 1, 1) if year else None
            except (TypeError, ValueError):
                born = None
        age = age_of(born)
        sexo = str(row[idx["Sexo"]] or "").strip()
        if sexo.lower().startswith("h"):
            sexo = "H"
        elif sexo.lower().startswith("m"):
            sexo = "M"
        else:
            sexo = sexo[:1].upper() if sexo else ""
        dni = str(row[idx["Identificador"]] or "").strip() if "Identificador" in idx else ""
        tipo_doc = str(row[idx["Tipo documento"]] or "").strip() if "Tipo documento" in idx else ""
        if dni in {"", "-", "—", "."} or tipo_doc.lower() == "ninguna":
            dni = ""
            tipo_doc = ""
        person = {
            "nombre": " ".join(
                str(row[idx[h]] or "").strip()
                for h in ("Nombre", "Apellido 1", "Apellido 2")
                if str(row[idx[h]] or "").strip()
            ),
            "sexo": sexo,
            "edad": age,
            "anio": born.year if born else None,
            "nacionalidad": str(row[idx["Nacionalidad"]] or "").strip() or None,
            "dni": dni or None,
            "tipo_doc": tipo_doc or None,
        }
        for key in address_keys(direccion):
            by[key].append(person)
        count += 1
        if count % 25000 == 0:
            print(f"  {count} personas", flush=True)
    wb.close()
    print(f"Personas leídas: {count}  claves: {len(by)}", flush=True)
    return by


def pick_people(
    by_addr: dict[str, list[dict]],
    row: dict,
    claimed: set[tuple] | None = None,
) -> list[dict]:
    """Asigna personas a UNA vivienda: claves específicas primero; sin reutilizar claimed."""
    claimed = claimed if claimed is not None else set()
    piso = (row.get("PISO") or "").strip()
    puerta = (row.get("PUERTA") or "").strip()
    keys = address_keys(
        row.get("DireccionOriginal") or "",
        row.get("VIA_NUMERO") or "",
        row.get("PORTAL") or "",
        piso,
        puerta,
    )
    # Con planta/puerta no usar claves de solo portal (evita copiar vecinos a todos los pisos)
    if piso or puerta:
        keys = [k for k in keys if not is_portal_level_key(k)]

    for key in keys:
        if key not in by_addr:
            continue
        seen = set()
        out = []
        for person in by_addr[key]:
            sig = person_sig(person)
            if not sig[0] or sig in seen or sig in claimed:
                continue
            seen.add(sig)
            out.append(person)
        if out:
            for person in out:
                claimed.add(person_sig(person))
            return out
    return []


def aggregates(people: list[dict], fallback_hab: int) -> dict:
    h = sum(1 for p in people if p.get("sexo") == "H")
    m = sum(1 for p in people if p.get("sexo") == "M")
    ages = [p["edad"] for p in people if p.get("edad") is not None]
    bands = [0] * len(AGE_BANDS)
    for age in ages:
        i = band_index(age)
        if i is not None:
            bands[i] += 1
    n = len(people) if people else fallback_hab
    return {
        "residentes": n,
        "hombres": h,
        "mujeres": m,
        "edad_media": round(sum(ages) / len(ages), 1) if ages else None,
        "bandas": bands,
    }


def load_coord_index():
    """Índice dirección/ref → (lon, lat, z) desde el geojson actual."""
    by_dir: dict[str, tuple[float, float, float]] = {}
    by_ref: dict[str, tuple[float, float, float]] = {}
    by_rc: dict[str, list[tuple[float, float, float]]] = defaultdict(list)
    if GEO_IN.exists():
        geo = json.loads(GEO_IN.read_text(encoding="utf-8"))
        for feat in geo.get("features") or []:
            coords = (feat.get("geometry") or {}).get("coordinates") or []
            if len(coords) < 2:
                continue
            lon, lat = float(coords[0]), float(coords[1])
            z = float(coords[2]) if len(coords) > 2 else 1.5
            props = feat.get("properties") or {}
            d = norm_key(props.get("direccion") or "")
            if d:
                by_dir[d] = (lon, lat, z)
                by_dir[re.sub(r"\b0+(\d)", r"\1", d)] = (lon, lat, z)
            ref = (props.get("ref") or "").strip().upper()
            rc = (props.get("rc") or "").strip().upper()
            if ref:
                by_ref[ref] = (lon, lat, z)
            if rc:
                by_rc[rc].append((lon, lat, z))
    cache = {}
    if CACHE.exists():
        cache = json.loads(CACHE.read_text(encoding="utf-8"))
    manual = {}
    if MANUAL_XY.exists():
        raw = json.loads(MANUAL_XY.read_text(encoding="utf-8"))
        for k, v in (raw or {}).items():
            if not v or len(v) < 2:
                continue
            nk = norm_key(k)
            manual[nk] = (float(v[0]), float(v[1]), 1.5)
            manual[re.sub(r"\b0+(\d)", r"\1", nk)] = manual[nk]
    return by_dir, by_ref, by_rc, cache, manual


def resolve_xy(row, by_dir, by_ref, by_rc, cache, manual=None, portales_xy=None):
    ref = (row.get("ReferenciaCatastral") or "").strip().upper()
    rc = (row.get("REFCAT_14") or ref[:14]).strip().upper()
    keys = address_keys(
        row.get("DireccionOriginal") or "",
        row.get("VIA_NUMERO") or "",
        row.get("PORTAL") or "",
        row.get("PISO") or "",
        row.get("PUERTA") or "",
    )
    manual = manual or {}
    portales_xy = portales_xy or {}
    z = z_for(row.get("PISO") or "")
    for key in keys:
        if key in manual:
            return manual[key]
    # Portal municipal por calle+número (fachada correcta en parcelas con varias calles)
    for sn in street_num_keys(row.get("VIA_NUMERO") or "", row.get("DireccionOriginal") or ""):
        if sn in portales_xy:
            lon, lat = portales_xy[sn]
            return lon, lat, z
        if rc and f"{rc}|{sn}" in portales_xy:
            lon, lat = portales_xy[f"{rc}|{sn}"]
            return lon, lat, z
    for key in keys:
        if key in by_dir:
            return by_dir[key]
    if ref and ref in by_ref:
        return by_ref[ref]
    if rc and rc in by_rc and by_rc[rc]:
        pts = by_rc[rc]
        lon = sum(p[0] for p in pts) / len(pts)
        lat = sum(p[1] for p in pts) / len(pts)
        return lon, lat, z
    if rc and rc in cache and len(cache[rc]) == 2:
        return float(cache[rc][0]), float(cache[rc][1]), z
    return None


def portal_xy_for_via(via: str, row: dict, portales_xy: dict) -> tuple[float, float] | None:
    for sn in street_num_keys(via, row.get("DireccionOriginal") or ""):
        if sn in portales_xy:
            return portales_xy[sn]
        rc = (row.get("REFCAT_14") or "").strip().upper()
        if rc and f"{rc}|{sn}" in portales_xy:
            return portales_xy[f"{rc}|{sn}"]
    return None


def main():
    by_people = load_people_by_address()
    by_dir, by_ref, by_rc, cache, manual = load_coord_index()
    portales_xy = load_portales_xy()

    rows = []
    with PADRON.open(encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            match = (r.get("match") or "").strip()
            if match.startswith("excluded"):
                continue
            rows.append(r)
    print(f"Filas padrón (sin excluded): {len(rows)}", flush=True)

    features = []
    personas = {}
    edificios_map = {}
    skipped = 0
    matched_people = 0
    vid = 0
    claimed_people: set[tuple] = set()

    # Agrupar por parcela para tipología
    by_parcel = defaultdict(list)
    orphans = []
    for r in rows:
        rc = (r.get("REFCAT_14") or (r.get("ReferenciaCatastral") or "")[:14] or "").strip().upper()
        if len(rc) >= 14:
            by_parcel[rc].append(r)
        else:
            orphans.append(r)

    def emit_row(r, lon, lat, z, tipo_force=None):
        nonlocal vid, matched_people
        vid += 1
        try:
            hab = int((r.get("HABITANTES") or "0").strip() or 0)
        except ValueError:
            hab = 0
        people = pick_people(by_people, r, claimed_people)
        if people:
            matched_people += len(people)
        agg = aggregates(people, hab)
        ref = (r.get("ReferenciaCatastral") or "").strip().upper()
        rc = (r.get("REFCAT_14") or ref[:14]).strip().upper()
        via = (r.get("VIA_NUMERO") or r.get("DireccionOriginal") or "Sin dirección").strip()
        match = (r.get("match") or "").strip()
        incidencia = ""
        if match == "ok_canada_parcela":
            incidencia = "Posición aproximada por parcela Cañada (sin ref. catastral)."
        elif match == "not_in_catastro":
            incidencia = "Sin ref. catastral. Situada manualmente."
        unica = tipo_force == "unifamiliar" or (tipo_force is None and True)
        tipo = tipo_force or "unifamiliar"
        props = {
            "id": vid,
            "rc": rc,
            "ref": ref,
            "direccion": via,
            "planta": planta_txt(r.get("PISO") or "", tipo == "unifamiliar"),
            "letra": (r.get("PUERTA") or r.get("PORTAL") or "").strip(),
            "portal": (r.get("PORTAL") or "").strip(),
            "piso": (r.get("PISO") or "").strip(),
            "residentes": agg["residentes"],
            "hombres": agg["hombres"],
            "mujeres": agg["mujeres"],
            "edad_media": agg["edad_media"],
            "b0": agg["bandas"][0],
            "b1": agg["bandas"][1],
            "b2": agg["bandas"][2],
            "b3": agg["bandas"][3],
            "b4": agg["bandas"][4],
            "tipo": tipo,
            "incidencia": incidencia,
            "vid": vid,
        }
        features.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [round(lon, 7), round(lat, 7), round(z, 2)]},
                "properties": props,
            }
        )
        if people:
            personas[str(vid)] = people
        return props, agg

    def offset_xy(lon: float, lat: float, index: int, total: int, meters: float = 14.0):
        if total <= 1:
            return lon, lat
        angle = (2 * math.pi * index) / total - math.pi / 2
        cos = math.cos(lat * math.pi / 180) or 1.0
        return (
            lon + (meters * math.sin(angle)) / (111320 * cos),
            lat + (meters * math.cos(angle)) / 110540,
        )

    # Por parcela: tipología y coordenadas por fachada (VIA_NUMERO / portal municipal)
    for rc, group in by_parcel.items():
        placed = []
        for r in group:
            xy = resolve_xy(r, by_dir, by_ref, by_rc, cache, manual, portales_xy)
            if not xy:
                skipped += 1
                continue
            placed.append((r, xy))
        if not placed:
            continue
        mean_lon = sum(xy[0] for _, xy in placed) / len(placed)
        mean_lat = sum(xy[1] for _, xy in placed) / len(placed)
        by_via: dict[str, list] = defaultdict(list)
        for r, xy in placed:
            via = (r.get("VIA_NUMERO") or "").strip() or "Sin dirección"
            by_via[via].append((r, xy))
        via_names = list(by_via.keys())
        # ¿Cada fachada tiene portal propio? Si sí, no hace falta el offset circular.
        via_portal: dict[str, tuple[float, float]] = {}
        for via in via_names:
            sample = by_via[via][0][0]
            pt = portal_xy_for_via(via, sample, portales_xy)
            if pt:
                via_portal[via] = pt
        total_res = 0
        for vi, via in enumerate(via_names):
            items = by_via[via]
            tipo = "unifamiliar" if len(items) == 1 else "plurifamiliar"
            if via in via_portal:
                lon, lat = via_portal[via]
            elif len(via_portal) == len(via_names):
                lon = sum(xy[0] for _, xy in items) / len(items)
                lat = sum(xy[1] for _, xy in items) / len(items)
            else:
                lon, lat = offset_xy(mean_lon, mean_lat, vi, len(via_names))
            for r, xy in items:
                props, agg = emit_row(r, lon, lat, xy[2], tipo_force=tipo)
                total_res += agg["residentes"]
            sample = items[0][0]
            ref = (sample.get("ReferenciaCatastral") or "").strip().upper()
            key = rc if len(via_names) == 1 else f"{rc}|{via}"
            edificios_map[key] = {
                "rc": rc,
                "ref": ref if len(items) == 1 else rc,
                "lon": round(lon, 7),
                "lat": round(lat, 7),
                "direccion": via,
                "viviendas": len(items),
                "residentes": sum(
                    int(str((row.get("HABITANTES") or "0")).strip() or 0)
                    for row, _ in items
                ),
                "tipo": tipo,
                "planta": planta_txt(sample.get("PISO") or "", tipo == "unifamiliar")
                if tipo == "unifamiliar"
                else "",
                "letra": (sample.get("PUERTA") or "").strip() if tipo == "unifamiliar" else "",
            }

    for r in orphans:
        xy = resolve_xy(r, by_dir, by_ref, by_rc, cache, manual, portales_xy)
        if not xy:
            skipped += 1
            continue
        props, agg = emit_row(r, xy[0], xy[1], xy[2], tipo_force="unifamiliar")
        key = f"orphan-{props['vid']}"
        edificios_map[key] = {
            "rc": "",
            "ref": props["ref"],
            "lon": round(xy[0], 7),
            "lat": round(xy[1], 7),
            "direccion": props["direccion"],
            "viviendas": 1,
            "residentes": agg["residentes"],
            "tipo": "unifamiliar",
            "planta": props["planta"],
            "letra": props["letra"],
            "incidencia": props["incidencia"],
        }

    edificios = list(edificios_map.values())
    lons = [e["lon"] for e in edificios]
    lats = [e["lat"] for e in edificios]
    center = [
        round(sum(lons) / len(lons), 6),
        round(sum(lats) / len(lats), 6),
    ] if edificios else [-3.53, 40.35]

    geojson = {"type": "FeatureCollection", "features": features}
    (DATA / "viviendas.geojson").write_text(
        json.dumps(geojson, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    (DATA / "edificios.json").write_text(
        json.dumps(edificios, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    (DATA / "personas.json").write_text(
        json.dumps(personas, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    meta = {
        "viviendas": len(features),
        "edificios": len(edificios),
        "personas_indexadas": len(personas),
        "personas_vinculadas": matched_people,
        "sin_coordenada": skipped,
        "center": center,
        "age_bands": AGE_LABELS,
        "fecha_padron": "2026-10-08",
        "fuente": "padron_refcat+habitantes",
    }
    (DATA / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"OK viviendas={len(features)} edificios={len(edificios)} "
        f"personas_vids={len(personas)} vinculadas={matched_people} sin_xy={skipped}",
        flush=True,
    )


if __name__ == "__main__":
    main()
